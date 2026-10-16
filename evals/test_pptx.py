"""Test cho PptxBuilder (pptx.py).

Chạy:  python3 -m pytest test_pptx.py -v   (hoặc: python3 test_pptx.py)
"""
import io
import sys
import types
import zipfile
import importlib.util
import xml.etree.ElementTree as ET
from pathlib import Path

import pytest

HERE = Path(__file__).parent


def load_builder(mod_name="pptx_builder_mod"):
    """Load pptx.py theo đường dẫn file, tránh đụng tên package 'pptx'."""
    spec = importlib.util.spec_from_file_location(
        mod_name, HERE.parent / "agent" / "src" / "documents" / "pptx.py")
    mod = importlib.util.module_from_spec(spec)
    sys.modules[mod_name] = mod
    spec.loader.exec_module(mod)
    return mod


# ---------------------------------------------------------------------------
# Stub giả lập python-pptx để test nhánh native mà không cần cài lib thật
# ---------------------------------------------------------------------------
def install_pptx_stub(log):
    pptx = types.ModuleType("pptx")
    util = types.ModuleType("pptx.util")
    util.Inches = lambda v: int(v * 914400)
    util.Pt = lambda v: int(v * 12700)

    class _Run:
        def __init__(self):
            self.font = types.SimpleNamespace(bold=False)

    class _Para:
        def __init__(self):
            self.text = ""
            self.level = 0
            self.runs = [_Run()]

    class _TextFrame:
        def __init__(self):
            self.paragraphs = [_Para()]

        def clear(self):
            self.paragraphs = [_Para()]

        def add_paragraph(self):
            p = _Para()
            self.paragraphs.append(p)
            return p

    class _Cell:
        def __init__(self):
            self.text = ""
            self._tf = _TextFrame()

        @property
        def text_frame(self):
            return self._tf

    class _Table:
        def __init__(self, nr, nc):
            self._cells = [[_Cell() for _ in range(nc)] for _ in range(nr)]
            self.columns = [types.SimpleNamespace(width=0) for _ in range(nc)]

        def cell(self, i, j):
            return self._cells[i][j]

    class _Shapes:
        def __init__(self):
            self.title = types.SimpleNamespace(text="")

        def add_table(self, nr, nc, left, top, width, height):
            log.append(("add_table", nr, nc))
            return types.SimpleNamespace(table=_Table(nr, nc))

    class _Slide:
        def __init__(self, layout):
            self.shapes = _Shapes()
            self._ph = {1: types.SimpleNamespace(text="", text_frame=_TextFrame())}
            log.append(("add_slide", layout))

        @property
        def placeholders(self):
            return self._ph

    class _Slides:
        def __init__(self):
            self.items = []

        def add_slide(self, layout):
            sl = _Slide(layout)
            self.items.append(sl)
            return sl

    class Presentation:
        def __init__(self):
            self.slide_layouts = [0, 1, 2, 3, 4, 5, 6]
            self.slides = _Slides()
            log.append(("Presentation",))

        def save(self, bio):
            bio.write(b"FAKEPPTX:" + str(len(self.slides.items)).encode())

    pptx.Presentation = Presentation
    sys.modules["pptx"] = pptx
    sys.modules["pptx.util"] = util


@pytest.fixture()
def pptx_stubbed():
    saved = {k: sys.modules.get(k) for k in ("pptx", "pptx.util")}
    log = []
    install_pptx_stub(log)
    yield log
    for k, v in saved.items():
        if v is None:
            sys.modules.pop(k, None)
        else:
            sys.modules[k] = v


# ---------------------------------------------------------------------------
# Test fallback (không có python-pptx)
# ---------------------------------------------------------------------------
def test_fallback_three_slides_save(tmp_path):
    mod = load_builder("pptx_fb1")
    b = mod.PptxBuilder(force_fallback=True)
    assert b.available is False
    (b.add_title_slide("Báo cáo Q3", "Boom team")
      .add_bullet_slide("Điểm chính", ["Tăng trưởng 20%", "Mở rộng 3 kênh"])
      .add_table_slide("Số liệu", ["Kênh", "Doanh thu"],
                       [["Online", "1.2 tỷ"], ["Retail", "800 triệu"]]))
    assert b.slide_count == 3

    out = b.save(str(tmp_path / "deck.pptx"))
    data = Path(out).read_bytes()
    assert len(data) > 0, "file pptx fallback không được rỗng"

    # là ZIP hợp lệ, đủ part chuẩn OOXML
    assert zipfile.is_zipfile(io.BytesIO(data))
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        names = z.namelist()
        assert "[Content_Types].xml" in names
        assert "ppt/presentation.xml" in names
        assert "ppt/theme/theme1.xml" in names
        slides = sorted(n for n in names
                        if n.startswith("ppt/slides/slide") and n.endswith(".xml"))
        assert len(slides) == 3, f"expect 3 slides, got {slides}"
        # mọi part XML đều well-formed
        for n in names:
            if n.endswith(".xml"):
                ET.fromstring(z.read(n))
        # nội dung các slide có mặt
        assert "Báo cáo Q3" in z.read("ppt/slides/slide1.xml").decode("utf-8")
        assert "Tăng trưởng 20%" in z.read("ppt/slides/slide2.xml").decode("utf-8")
        assert "Doanh thu" in z.read("ppt/slides/slide3.xml").decode("utf-8")


def test_fallback_auto_when_lib_missing():
    """Không ép fallback: máy thiếu python-pptx thì tự rơi vào fallback, không crash."""
    mod = load_builder("pptx_fb2")
    try:
        from pptx import Presentation as _real  # noqa  (lib thật, không phải builder local)
        if "pptx_builder" in (_real.__module__ or ""):
            raise ImportError
    except ImportError:
        pass
    else:
        pytest.skip("máy này đã có python-pptx, bỏ qua test auto-fallback")
    b = mod.PptxBuilder()
    assert b.available is False
    b.add_text_slide("Hello", "world")
    assert b.to_bytes()[:2] == b"PK"  # ZIP signature


def test_fluent_chaining_returns_self():
    mod = load_builder("pptx_fb3")
    b = mod.PptxBuilder(force_fallback=True)
    assert b.add_title_slide("t") is b
    assert b.add_bullet_slide("t", ["a"]) is b
    assert b.add_table_slide("t", ["h"], [["x"]]) is b
    assert b.add_text_slide("t", "x") is b
    assert b.save("/tmp/_chain_test.pptx") == "/tmp/_chain_test.pptx"


def test_to_bytes_nonempty():
    mod = load_builder("pptx_fb4")
    data = (mod.PptxBuilder(force_fallback=True)
            .add_title_slide("T", "S")
            .to_bytes())
    assert isinstance(data, bytes) and len(data) > 1000


# ---------------------------------------------------------------------------
# Test nhánh native (dùng stub thay cho python-pptx thật)
# ---------------------------------------------------------------------------
def test_native_path_with_stub(pptx_stubbed):
    log = pptx_stubbed
    mod = load_builder("pptx_native_stub")
    b = mod.PptxBuilder()
    assert b.available is True, "stub đã inject thì phải vào nhánh native"

    (b.add_title_slide("T", "S")
      .add_bullet_slide("B", ["a", "b"])
      .add_table_slide("Tab", ["H1", "H2"], [["a", "b"]])
      .add_text_slide("Tx", "hello"))
    data = b.to_bytes()
    assert data == b"FAKEPPTX:4"

    layouts = [e[1] for e in log if e[0] == "add_slide"]
    assert layouts == [0, 1, 5, 1], f"layout order sai: {layouts}"
    assert ("add_table", 2, 2) in log, "bảng 2x2 (1 header + 1 row) phải được tạo"


def test_native_readback_if_real_lib():
    """Nếu máy có python-pptx thật: build rồi đọc lại, verify đúng 3 slides."""
    try:
        from pptx import Presentation  # noqa
    except ImportError:
        pytest.skip("không có python-pptx thật")
    mod = load_builder("pptx_native_real")
    b = mod.PptxBuilder()
    assert b.available is True
    (b.add_title_slide("T", "S")
      .add_bullet_slide("B", ["a"])
      .add_table_slide("Tab", ["H"], [["x"]]))
    prs = Presentation(io.BytesIO(b.to_bytes()))
    assert len(prs.slides) == 3


if __name__ == "__main__":
    sys.exit(pytest.main([__file__, "-v"]))
