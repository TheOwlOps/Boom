#!/usr/bin/env python3
"""Test XlsxBuilder: tạo 2 sheets, verify công thức/data, to_bytes, fallback."""
import importlib.util
import sys
from unittest import mock

import pytest

import xlsx as xlsx_mod
from xlsx import XlsxBuilder


@pytest.fixture()
def wb():
    """Builder mẫu: sheet1 có data + công thức SUM, sheet2 có data."""
    b = XlsxBuilder()
    (b.add_sheet("Sales")
      .set_cell("Sales", 1, 1, "Tháng")
      .set_cell("Sales", 1, 2, "Doanh thu")
      .set_cell("Sales", 2, 1, "T1")
      .set_cell("Sales", 2, 2, 100)
      .set_cell("Sales", 3, 1, "T2")
      .set_cell("Sales", 3, 2, 200)
      .set_formula("Sales", 4, 2, "=SUM(B2:B3)")
      .set_formula("Sales", 5, 2, "SUM(B2:B3)")  # không có '=' -> tự thêm
      .style_cell("Sales", 1, 1, bold=True, color="FF0000", border=True)
      .auto_width("Sales")
      .add_chart("Sales", "bar", "B2:B3", "Doanh thu")
      .add_sheet("Notes")
      .set_cell("Notes", 1, 1, "Ghi chú")
      .set_cell("Notes", 2, 1, "Báo cáo tháng"))
    return b


@pytest.mark.skipif(not xlsx_mod._HAS_OPENPYXL, reason="cần openpyxl")
def test_save_and_readback(tmp_path, wb):
    path = str(tmp_path / "out.xlsx")
    assert wb.save(path) == path
    from openpyxl import load_workbook
    book = load_workbook(path)
    assert book.sheetnames == ["Sales", "Notes"]

    ws = book["Sales"]
    # data nguyên vẹn
    assert ws.cell(1, 1).value == "Tháng"
    assert ws.cell(2, 2).value == 100
    assert ws.cell(3, 2).value == 200
    # công thức giữ nguyên chuỗi
    assert ws.cell(4, 2).value == "=SUM(B2:B3)"
    assert ws.cell(5, 2).value == "=SUM(B2:B3)"
    # style đã áp
    assert ws.cell(1, 1).font.bold is True
    # chart đã thêm
    assert len(ws._charts) == 1

    ws2 = book["Notes"]
    assert ws2.cell(1, 1).value == "Ghi chú"
    assert ws2.cell(2, 1).value == "Báo cáo tháng"


def test_to_bytes_not_empty(wb):
    data = wb.to_bytes()
    assert isinstance(data, bytes)
    assert len(data) > 0
    if xlsx_mod._HAS_OPENPYXL:
        # file xlsx là zip -> bắt đầu bằng PK
        assert data[:2] == b"PK"


def test_fluent_chaining_returns_self():
    b = XlsxBuilder()
    assert b.add_sheet("S") is b
    assert b.set_cell("S", 1, 1, "x") is b
    assert b.set_formula("S", 1, 2, "=1+1") is b
    assert b.style_cell("S", 1, 1, bold=True) is b
    assert b.auto_width("S") is b
    assert b.add_chart("S") is b


def test_invalid_chart_type_raises():
    pytest.importorskip("openpyxl")
    with pytest.raises(ValueError, match="chart_type không hợp lệ"):
        XlsxBuilder().add_chart("S", chart_type="radar")


def _load_fallback_module():
    """Load xlsx.py với openpyxl bị chặn -> mô phỏng máy thiếu lib."""
    with mock.patch.dict(sys.modules, {"openpyxl": None}):
        spec = importlib.util.spec_from_file_location(
            "xlsx_fallback", xlsx_mod.__file__)
        mod = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(mod)
    return mod


def test_fallback_no_crash(tmp_path):
    """Mock thiếu openpyxl: mọi API vẫn chạy, không crash."""
    mod = _load_fallback_module()
    assert mod._HAS_OPENPYXL is False
    b = mod.XlsxBuilder()
    assert b._has_openpyxl is False
    # Fluent API đầy đủ — không được raise
    (b.add_sheet("A")
      .set_cell("A", 1, 1, "x")
      .set_formula("A", 1, 2, "=SUM(A1:A1)")
      .style_cell("A", 1, 1, bold=True, color="FF0000", border=True)
      .auto_width("A")
      .add_chart("A", "bar", "B1:B2", "T"))
    out = str(tmp_path / "fallback.xlsx")
    saved = b.save(out)  # fallback: lưu CSV
    assert saved.endswith(".csv")
    with open(saved, encoding="utf-8") as fh:
        assert "x" in fh.read()
    assert len(b.to_bytes()) > 0
