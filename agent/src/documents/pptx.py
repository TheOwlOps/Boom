"""PptxBuilder — tạo file PowerPoint (.pptx) bằng python-pptx nếu có,
fallback sang writer thuần Python (zip + XML) nếu thiếu lib.

Cả hai đường đều cho ra file .pptx hợp lệ, API fluent (method chaining).
"""
from __future__ import annotations

import io
import zipfile
from typing import List
from xml.sax.saxutils import escape as _xml_escape

# ---------------------------------------------------------------------------
# Fallback writer thuần Python: dựng file .pptx tối giản bằng tay.
#
# .pptx thực chất là 1 file ZIP chứa các part XML theo chuẩn OOXML.
# Để PowerPoint/LibreOffice mở được, tối thiểu cần:
#   [Content_Types].xml, _rels/.rels, ppt/presentation.xml (+ rels của nó),
#   ppt/slideMasters/slideMaster1.xml, ppt/slideLayouts/slideLayout1.xml,
#   ppt/theme/theme1.xml, và mỗi slide: ppt/slides/slideN.xml (+ rels).
# ---------------------------------------------------------------------------

_EMU_PER_INCH = 914400
_SLIDE_W = 10 * _EMU_PER_INCH
_SLIDE_H = int(7.5 * _EMU_PER_INCH)

_NS = (
    'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
    'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
    'xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"'
)
_REL_NS = 'xmlns="http://schemas.openxmlformats.org/package/2006/relationships"'
_CT_NS = 'xmlns="http://schemas.openxmlformats.org/package/2006/content-types"'

_REL_OFFICEDOC = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument"
_REL_SLIDE = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/slide"
_REL_SLIDEMASTER = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/slideMaster"
_REL_SLIDELAYOUT = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/slideLayout"
_REL_THEME = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/theme"


def _para_xml(text: str, level: int = 0, bullet: bool = False,
              size: int = 1800, bold: bool = False, align: str = "") -> str:
    """Render 1 đoạn văn DrawingML (<a:p>). size tính bằng 1/100 pt."""
    t = _xml_escape(text)
    ppr_attrs = f' marL="0" indent="0"{f" algn=\"{align}\"" if align else ""}'
    if bullet:
        ppr = f"<a:pPr{ppr_attrs}><a:buChar char=\"\u2022\"/></a:pPr>"
    else:
        ppr = f"<a:pPr{ppr_attrs}/>"
    if t:
        rpr = f'<a:rPr lang="en-US" sz="{size}"{" b=\"1\"" if bold else ""}/>'
        return f"<a:p>{ppr}<a:r>{rpr}<a:t>{t}</a:t></a:r></a:p>"
    return f"<a:p>{ppr}<a:endParaRPr lang=\"en-US\"/></a:p>"


def _shape_xml(sid: int, name: str, x: int, y: int, cx: int, cy: int,
               paras: list, ph: str | None = None) -> str:
    """Render 1 textbox shape (<p:sp>) đặt tại (x, y), kích thước (cx, cy)."""
    nvpr = f'<p:nvPr><p:ph type="{ph}"/></p:nvPr>' if ph else "<p:nvPr/>"
    return (
        f'<p:sp><p:nvSpPr><p:cNvPr id="{sid}" name="{_xml_escape(name)}"/>'
        f'<p:cNvSpPr><a:spLocks noGrp="1"/></p:cNvSpPr>{nvpr}</p:nvSpPr>'
        f'<p:spPr><a:xfrm><a:off x="{x}" y="{y}"/><a:ext cx="{cx}" cy="{cy}"/>'
        f'</a:xfrm><a:prstGeom prst="rect"><a:avLst/></a:prstGeom></p:spPr>'
        f'<p:txBody><a:bodyPr wrap="square"/><a:lstStyle/>{"".join(paras)}</p:txBody></p:sp>'
    )


def _slide_xml(shapes: str) -> str:
    return (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        f'<p:sld {_NS}>'
        '<p:cSld><p:spTree>'
        '<p:nvGrpSpPr><p:cNvPr id="1" name=""/><p:cNvGrpSpPr/><p:nvPr/></p:nvGrpSpPr>'
        '<p:grpSpPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="0" cy="0"/>'
        '<a:chOff x="0" y="0"/><a:chExt cx="0" cy="0"/></a:xfrm></p:grpSpPr>'
        f'{shapes}'
        '</p:spTree></p:cSld>'
        '<p:clrMapOvr><a:overrideClrMapping bg1="lt1" tx1="dk1" bg2="lt2" tx2="dk2" '
        'accent1="accent1" accent2="accent2" accent3="accent3" accent4="accent4" '
        'accent5="accent5" accent6="accent6" hlink="hlink" folHlink="folHlink"/>'
        '</p:clrMapOvr></p:sld>'
    )


def _title_shape(sid: int, title: str, x: int, y: int, cx: int, cy: int,
                 size: int = 3200, align: str = "ctr") -> str:
    return _shape_xml(sid, "Title", x, y, cx, cy,
                      [_para_xml(title, size=size, bold=True, align=align)], ph="title")


def _slide_shapes_for_op(op: tuple) -> str:
    """Chuyển 1 op (định nghĩa slide) thành XML shapes cho fallback writer."""
    kind = op[0]
    inch = _EMU_PER_INCH
    tx, ty, tcx, tcy = int(0.5 * inch), int(0.25 * inch), int(9 * inch), int(1 * inch)
    bx, by, bcx, bcy = int(0.5 * inch), int(1.5 * inch), int(9 * inch), int(5.5 * inch)

    if kind == "title":
        _, title, subtitle = op
        s = _title_shape(2, title, int(0.5 * inch), int(1.5 * inch),
                         int(9 * inch), int(2 * inch), size=4000)
        if subtitle:
            s += _shape_xml(3, "Subtitle", int(1 * inch), int(3.75 * inch),
                            int(8 * inch), int(1.5 * inch),
                            [_para_xml(subtitle, size=2000, align="ctr")], ph="body")
        return s

    if kind == "bullet":
        _, title, bullets = op
        paras = [_para_xml(str(b), bullet=True, size=1800) for b in bullets] or [_para_xml("")]
        return (_title_shape(2, title, tx, ty, tcx, tcy, size=2800, align="l") +
                _shape_xml(3, "Content", bx, by, bcx, bcy, paras, ph="body"))

    if kind == "table":
        _, title, headers, rows = op
        # Fallback: bảng vẽ thành text căn cột bằng ký tự ' | ' (không đẹp bằng
        # bảng thật của python-pptx nhưng vẫn đọc được đầy đủ dữ liệu).
        cols = len(headers)
        widths = [0] * cols
        all_rows = [headers] + rows
        for r in all_rows:
            for j in range(cols):
                widths[j] = max(widths[j], len(str(r[j])) if j < len(r) else 0)
        def fmt(r):
            return " | ".join(str(r[j]).ljust(widths[j]) if j < len(r) else "".ljust(widths[j])
                              for j in range(cols))
        paras = [_para_xml(fmt(headers), size=1600, bold=True)]
        paras.append(_para_xml("-+-".join("-" * w for w in widths), size=1600))
        for r in rows:
            paras.append(_para_xml(fmt(r), size=1600))
        return (_title_shape(2, title, tx, ty, tcx, tcy, size=2800, align="l") +
                _shape_xml(3, "Content", bx, by, bcx, bcy, paras, ph="body"))

    if kind == "text":
        _, title, text = op
        paras = [_para_xml(line, size=1800) for line in str(text).splitlines()] or [_para_xml("")]
        return (_title_shape(2, title, tx, ty, tcx, tcy, size=2800, align="l") +
                _shape_xml(3, "Content", bx, by, bcx, bcy, paras, ph="body"))

    raise ValueError(f"Unknown slide op: {kind}")


_THEME_XML = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<a:theme xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" name="Boom Minimal">
<a:themeElements><a:clrScheme name="Boom">
<a:dk1><a:sysClr val="windowText" lastClr="000000"/></a:dk1>
<a:lt1><a:sysClr val="window" lastClr="FFFFFF"/></a:lt1>
<a:dk2><a:srgbClr val="1F2937"/></a:dk2>
<a:lt2><a:srgbClr val="EEECE1"/></a:lt2>
<a:accent1><a:srgbClr val="4472C4"/></a:accent1>
<a:accent2><a:srgbClr val="ED7D31"/></a:accent2>
<a:accent3><a:srgbClr val="A5A5A5"/></a:accent3>
<a:accent4><a:srgbClr val="FFC000"/></a:accent4>
<a:accent5><a:srgbClr val="5B9BD5"/></a:accent5>
<a:accent6><a:srgbClr val="70AD47"/></a:accent6>
<a:hlink><a:srgbClr val="0563C1"/></a:hlink>
<a:folHlink><a:srgbClr val="954F72"/></a:folHlink>
</a:clrScheme>
<a:fmtScheme name="Boom">
<a:fillStyleLst>
<a:solidFill><a:schemeClr val="phClr"/></a:solidFill>
<a:gradFill rotWithShape="0"><a:gsLst><a:gs pos="0"><a:schemeClr val="phClr"><a:tint val="50000"/><a:satMod val="300000"/></a:schemeClr></a:gs><a:gs pos="100000"><a:schemeClr val="phClr"><a:shade val="50000"/><a:satMod val="300000"/></a:schemeClr></a:gs></a:gsLst><a:lin ang="5400000" scaled="0"/></a:gradFill>
<a:gradFill rotWithShape="0"><a:gsLst><a:gs pos="0"><a:schemeClr val="phClr"><a:shade val="50000"/><a:satMod val="300000"/></a:schemeClr></a:gs><a:gs pos="100000"><a:schemeClr val="phClr"><a:tint val="50000"/><a:satMod val="300000"/></a:schemeClr></a:gs></a:gsLst><a:lin ang="5400000" scaled="0"/></a:gradFill>
</a:fillStyleLst>
<a:lnStyleLst>
<a:ln w="6350" cap="flat" cmpd="sng" algn="ctr"><a:solidFill><a:schemeClr val="phClr"/></a:solidFill><a:prstDash val="solid"/></a:ln>
<a:ln w="12700" cap="flat" cmpd="sng" algn="ctr"><a:solidFill><a:schemeClr val="phClr"/></a:solidFill><a:prstDash val="solid"/></a:ln>
<a:ln w="19050" cap="flat" cmpd="sng" algn="ctr"><a:solidFill><a:schemeClr val="phClr"/></a:solidFill><a:prstDash val="solid"/></a:ln>
</a:lnStyleLst>
<a:effectStyleLst>
<a:effectStyle><a:effectLst/></a:effectStyle>
<a:effectStyle><a:effectLst/></a:effectStyle>
<a:effectStyle><a:effectLst><a:outerShdw blurRad="57150" dist="38100" dir="5400000" algn="ctr" rotWithShape="0"><a:srgbClr val="000000"><a:alpha val="63000"/></a:srgbClr></a:outerShdw></a:effectLst></a:effectStyle>
</a:effectStyleLst>
<a:bgFillStyleLst>
<a:solidFill><a:schemeClr val="phClr"/></a:solidFill>
<a:solidFill><a:schemeClr val="phClr"/></a:solidFill>
<a:gradFill rotWithShape="0"><a:gsLst><a:gs pos="0"><a:schemeClr val="phClr"><a:tint val="50000"/><a:satMod val="300000"/></a:schemeClr></a:gs><a:gs pos="100000"><a:schemeClr val="phClr"><a:tint val="50000"/><a:satMod val="300000"/></a:schemeClr></a:gs></a:gsLst><a:lin ang="5400000" scaled="0"/></a:gradFill>
</a:bgFillStyleLst>
</a:fmtScheme></a:themeElements></a:theme>
"""


def _build_minimal_pptx(ops: list) -> bytes:
    """Dựng toàn bộ file .pptx (ZIP các part XML) từ danh sách ops."""
    n = len(ops)
    parts: list[tuple[str, str]] = []

    # [Content_Types].xml — khai báo MIME cho từng part
    overrides = [
        '<Override PartName="/ppt/presentation.xml" '
        'ContentType="application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml"/>',
        '<Override PartName="/ppt/slideMasters/slideMaster1.xml" '
        'ContentType="application/vnd.openxmlformats-officedocument.presentationml.slideMaster+xml"/>',
        '<Override PartName="/ppt/slideLayouts/slideLayout1.xml" '
        'ContentType="application/vnd.openxmlformats-officedocument.presentationml.slideLayout+xml"/>',
        '<Override PartName="/ppt/theme/theme1.xml" '
        'ContentType="application/vnd.openxmlformats-officedocument.presentationml.theme+xml"/>',
    ]
    for i in range(1, n + 1):
        overrides.append(
            f'<Override PartName="/ppt/slides/slide{i}.xml" '
            'ContentType="application/vnd.openxmlformats-officedocument.presentationml.slide+xml"/>')
    parts.append(("[Content_Types].xml",
                  '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                  f'<Types {_CT_NS}>'
                  '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
                  '<Default Extension="xml" ContentType="application/xml"/>'
                  + "".join(overrides) + "</Types>"))

    # _rels/.rels — trỏ package tới presentation chính
    parts.append(("_rels/.rels",
                  '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                  f'<Relationships {_REL_NS}>'
                  f'<Relationship Id="rId1" Type="{_REL_OFFICEDOC}" Target="ppt/presentation.xml"/>'
                  "</Relationships>"))

    # ppt/presentation.xml — danh sách slide + kích thước slide
    sld_ids = "".join(f'<p:sldId id="{256 + i}" r:id="rId{i + 2}"/>' for i in range(n))
    parts.append(("ppt/presentation.xml",
                  '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                  f'<p:presentation {_NS}>'
                  '<p:sldMasterIdLst><p:sldMasterId r:id="rId1"/></p:sldMasterIdLst>'
                  f'<p:sldIdLst>{sld_ids}</p:sldIdLst>'
                  f'<p:sldSz cx="{_SLIDE_W}" cy="{_SLIDE_H}"/>'
                  '<p:notesSz cx="6858000" cy="9144000"/>'
                  "</p:presentation>"))
    pres_rels = "".join(
        f'<Relationship Id="rId{i + 2}" Type="{_REL_SLIDE}" Target="slides/slide{i + 1}.xml"/>'
        for i in range(n))
    parts.append(("ppt/_rels/presentation.xml.rels",
                  '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                  f'<Relationships {_REL_NS}>'
                  f'<Relationship Id="rId1" Type="{_REL_SLIDEMASTER}" Target="slideMasters/slideMaster1.xml"/>'
                  f"{pres_rels}</Relationships>"))

    # Slide master tối giản
    parts.append(("ppt/slideMasters/slideMaster1.xml",
                  '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                  f'<p:sldMaster {_NS}>'
                  '<p:cSld><p:bg><p:bgPr/></p:bg><p:spTree>'
                  '<p:nvGrpSpPr><p:cNvPr id="1" name=""/><p:cNvGrpSpPr/><p:nvPr/></p:nvGrpSpPr>'
                  '<p:grpSpPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="0" cy="0"/>'
                  '<a:chOff x="0" y="0"/><a:chExt cx="0" cy="0"/></a:xfrm></p:grpSpPr>'
                  "</p:spTree></p:cSld>"
                  '<p:clrMap bg1="lt1" tx1="dk1" bg2="lt2" tx2="dk2" '
                  'accent1="accent1" accent2="accent2" accent3="accent3" '
                  'accent4="accent4" accent5="accent5" accent6="accent6" '
                  'hlink="hlink" folHlink="folHlink"/>'
                  '<p:sldLayoutIdLst><p:sldLayoutId r:id="rId1"/></p:sldLayoutIdLst>'
                  "<p:txStyles><p:titleStyle><a:lvl1pPr/></p:titleStyle>"
                  "<p:bodyStyle><a:lvl1pPr/></p:bodyStyle>"
                  "<p:otherStyle><a:lvl1pPr/></p:otherStyle></p:txStyles>"
                  "</p:sldMaster>"))
    parts.append(("ppt/slideMasters/_rels/slideMaster1.xml.rels",
                  '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                  f'<Relationships {_REL_NS}>'
                  f'<Relationship Id="rId1" Type="{_REL_SLIDELAYOUT}" Target="../slideLayouts/slideLayout1.xml"/>'
                  f'<Relationship Id="rId2" Type="{_REL_THEME}" Target="../theme/theme1.xml"/>'
                  "</Relationships>"))

    # Slide layout: chứa sẵn 2 placeholder title + body để slide kế thừa
    layout_shapes = (
        _shape_xml(2, "Title Placeholder", int(0.5 * _EMU_PER_INCH), int(0.25 * _EMU_PER_INCH),
                   int(9 * _EMU_PER_INCH), int(1 * _EMU_PER_INCH),
                   [_para_xml("Title", size=2800, bold=True)], ph="title") +
        _shape_xml(3, "Body Placeholder", int(0.5 * _EMU_PER_INCH), int(1.5 * _EMU_PER_INCH),
                   int(9 * _EMU_PER_INCH), int(5.5 * _EMU_PER_INCH),
                   [_para_xml("Content", bullet=True)], ph="body")
    )
    parts.append(("ppt/slideLayouts/slideLayout1.xml",
                  '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                  f'<p:sldLayout {_NS} type="titleAndContent" preserve="1">'
                  '<p:cSld name="Title and Content"><p:spTree>'
                  '<p:nvGrpSpPr><p:cNvPr id="1" name=""/><p:cNvGrpSpPr/><p:nvPr/></p:nvGrpSpPr>'
                  '<p:grpSpPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="0" cy="0"/>'
                  '<a:chOff x="0" y="0"/><a:chExt cx="0" cy="0"/></a:xfrm></p:grpSpPr>'
                  f"{layout_shapes}</p:spTree></p:cSld>"
                  '<p:clrMapOvr><a:overrideClrMapping bg1="lt1" tx1="dk1" bg2="lt2" tx2="dk2" '
                  'accent1="accent1" accent2="accent2" accent3="accent3" accent4="accent4" '
                  'accent5="accent5" accent6="accent6" hlink="hlink" folHlink="folHlink"/>'
                  '</p:clrMapOvr></p:sldLayout>'))
    parts.append(("ppt/slideLayouts/_rels/slideLayout1.xml.rels",
                  '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                  f'<Relationships {_REL_NS}>'
                  f'<Relationship Id="rId1" Type="{_REL_SLIDEMASTER}" Target="../slideMasters/slideMaster1.xml"/>'
                  "</Relationships>"))

    parts.append(("ppt/theme/theme1.xml", _THEME_XML))

    # Từng slide + rels của nó
    for i, op in enumerate(ops, start=1):
        parts.append((f"ppt/slides/slide{i}.xml", _slide_xml(_slide_shapes_for_op(op))))
        parts.append((f"ppt/slides/_rels/slide{i}.xml.rels",
                      '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                      f'<Relationships {_REL_NS}>'
                      f'<Relationship Id="rId1" Type="{_REL_SLIDELAYOUT}" Target="../slideLayouts/slideLayout1.xml"/>'
                      "</Relationships>"))

    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for name, data in parts:
            z.writestr(name, data.encode("utf-8"))
    return buf.getvalue()


class PptxBuilder:
    """Fluent builder tạo file .pptx.

    Dùng python-pptx khi có sẵn; nếu không, fallback sang writer thuần Python
    (không cần cài thêm gì). Không bao giờ crash vì thiếu lib.
    """

    def __init__(self, force_fallback: bool = False):
        self._ops: list[tuple] = []
        self._native = False
        if not force_fallback:
            try:
                from pptx import Presentation  # noqa: F401
                self._native = True
            except ImportError:
                self._native = False

    @property
    def available(self) -> bool:
        """True nếu đang dùng python-pptx (chế độ native), False nếu fallback."""
        return self._native

    @property
    def slide_count(self) -> int:
        return len(self._ops)

    # -- fluent API ---------------------------------------------------------
    def add_title_slide(self, title: str, subtitle: str = "") -> "PptxBuilder":
        self._ops.append(("title", title, subtitle))
        return self

    def add_bullet_slide(self, title: str, bullets: List[str]) -> "PptxBuilder":
        self._ops.append(("bullet", title, list(bullets)))
        return self

    def add_table_slide(self, title: str, headers: List[str],
                        rows: List[List]) -> "PptxBuilder":
        self._ops.append(("table", title, list(headers), [list(r) for r in rows]))
        return self

    def add_text_slide(self, title: str, text: str) -> "PptxBuilder":
        self._ops.append(("text", title, text))
        return self

    # -- output -------------------------------------------------------------
    def to_bytes(self) -> bytes:
        if self._native:
            buf = io.BytesIO()
            self._build_native().save(buf)
            return buf.getvalue()
        return _build_minimal_pptx(self._ops)

    def save(self, path: str) -> str:
        data = self.to_bytes()
        with open(path, "wb") as f:
            f.write(data)
        return path

    # -- native build (python-pptx) -----------------------------------------
    def _build_native(self) -> "object":
        from pptx import Presentation
        from pptx.util import Inches

        prs = Presentation()
        for op in self._ops:
            kind = op[0]
            if kind == "title":
                _, title, subtitle = op
                slide = prs.slides.add_slide(prs.slide_layouts[0])  # title slide
                slide.shapes.title.text = title
                try:
                    slide.placeholders[1].text = subtitle
                except (IndexError, KeyError):
                    pass
            elif kind == "bullet":
                _, title, bullets = op
                slide = prs.slides.add_slide(prs.slide_layouts[1])  # title + content
                slide.shapes.title.text = title
                tf = slide.placeholders[1].text_frame
                tf.clear()
                for i, b in enumerate(bullets):
                    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                    p.text = str(b)
                    p.level = 0
            elif kind == "table":
                _, title, headers, rows = op
                slide = prs.slides.add_slide(prs.slide_layouts[5])  # title only
                slide.shapes.title.text = title
                data = [[str(c) for c in headers]] + [[str(c) for c in r] for r in rows]
                nrows, ncols = len(data), max(len(headers), 1)
                width = Inches(9)
                tbl = slide.shapes.add_table(
                    nrows, ncols, Inches(0.5), Inches(1.5), width, Inches(0.5)).table
                for j in range(ncols):
                    tbl.columns[j].width = int(width / ncols)
                for i, row in enumerate(data):
                    for j in range(ncols):
                        cell = tbl.cell(i, j)
                        cell.text = row[j] if j < len(row) else ""
                        if i == 0:  # header: in đậm
                            for par in cell.text_frame.paragraphs:
                                for run in par.runs:
                                    run.font.bold = True
            elif kind == "text":
                _, title, text = op
                slide = prs.slides.add_slide(prs.slide_layouts[1])
                slide.shapes.title.text = title
                slide.placeholders[1].text = str(text)
            else:
                raise ValueError(f"Unknown slide op: {kind}")
        return prs
